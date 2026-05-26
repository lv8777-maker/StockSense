from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DARK = RGBColor(0x3C, 0x3C, 0x3B)
YELLOW = RGBColor(0xFD, 0xC8, 0x00)
LIGHT = RGBColor(0xA7, 0xA9, 0xAC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_bg(slide, color):
    bg = slide.shapes.add_shape(1, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def add_text(slide, left, top, width, height, text, *, size=18, bold=False, color=WHITE, align=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_accent_bar(slide, left, top, width, height, color=YELLOW):
    bar = slide.shapes.add_shape(1, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


# Slide 1 — Title
s = prs.slides.add_slide(BLANK)
add_bg(s, DARK)
add_accent_bar(s, Inches(0.6), Inches(3.0), Inches(0.25), Inches(1.6))
add_text(s, Inches(1.0), Inches(2.9), Inches(11), Inches(1.0),
         "Maverick Loyalty", size=54, bold=True, color=WHITE)
add_text(s, Inches(1.0), Inches(3.9), Inches(11), Inches(0.7),
         "Launch Checklist", size=32, bold=True, color=YELLOW)
add_text(s, Inches(1.0), Inches(4.7), Inches(11), Inches(0.5),
         "What's left to take CustomerConnect from dev build to live product",
         size=18, color=LIGHT)
add_text(s, Inches(0.6), Inches(7.0), Inches(12), Inches(0.3),
         "Prepared for Maverick Telecom", size=11, color=LIGHT)

# Slide 2 — Where we are today
s = prs.slides.add_slide(BLANK)
add_bg(s, WHITE)
add_accent_bar(s, 0, 0, SW, Inches(0.25), YELLOW)
add_text(s, Inches(0.6), Inches(0.5), Inches(12), Inches(0.7),
         "Where we are today", size=32, bold=True, color=DARK)
done = [
    "Email + phone sign-up with a single 6-digit verification code",
    "4-tier loyalty system (Starter, Explorer, Champion, Elite)",
    "Receipt OCR (Tesseract) and MTN tax-invoice processing",
    "Admin dashboard with role-based access control",
    "Security hardening: bcrypt, rate limits, CSRF, HTTP-only sessions",
]
remaining = [
    "Email delivery to real customers (SendGrid)",
    "Production database migration",
    "Account recovery for lost email/phone",
    "Legal pages and branding polish",
    "Real-device QA, deploy, and go-live",
]

def two_col(slide, left, top, w, h, title, items, title_color):
    add_text(slide, left, top, w, Inches(0.5), title, size=20, bold=True, color=title_color)
    y = top + Inches(0.6)
    for it in items:
        dot = slide.shapes.add_shape(9, left, y + Inches(0.08), Inches(0.18), Inches(0.18))
        dot.fill.solid()
        dot.fill.fore_color.rgb = title_color
        dot.line.fill.background()
        add_text(slide, left + Inches(0.35), y, w - Inches(0.35), Inches(0.6), it, size=14, color=DARK)
        y += Inches(0.55)

two_col(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(5.5),
        "Done", done, RGBColor(0x2E, 0x7D, 0x32))
two_col(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(5.5),
        "Still to do", remaining, RGBColor(0xC6, 0x28, 0x28))

# Slide 3 — How to read this deck
s = prs.slides.add_slide(BLANK)
add_bg(s, DARK)
add_text(s, Inches(0.6), Inches(0.6), Inches(12), Inches(0.8),
         "The 10 steps to launch", size=32, bold=True, color=YELLOW)
add_text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.6),
         "One slide per step. Each step says what it is, why it matters, and what 'done' looks like.",
         size=16, color=LIGHT)

steps_overview = [
    ("1", "Email delivery"),         ("2", "Production database"),
    ("3", "Account recovery"),       ("4", "Legal & content"),
    ("5", "Branding polish"),        ("6", "Real-device QA"),
    ("7", "Admin & catalogue"),      ("8", "Deploy & domain"),
    ("9", "Monitoring & backups"),   ("10", "Go-live announcement"),
]
col_w = Inches(2.4)
row_h = Inches(1.0)
for i, (num, name) in enumerate(steps_overview):
    col = i % 5
    row = i // 5
    x = Inches(0.6) + col * (col_w + Inches(0.1))
    y = Inches(2.5) + row * (row_h + Inches(0.3))
    box = s.shapes.add_shape(5, x, y, col_w, row_h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = YELLOW
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.margin_left = tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f"Step {num}"
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = YELLOW
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = name
    r2.font.size = Pt(16); r2.font.bold = True; r2.font.color.rgb = DARK

# Helper for step slides
def step_slide(num, title, what, why, done_text):
    s = prs.slides.add_slide(BLANK)
    add_bg(s, WHITE)
    # Side band
    band = s.shapes.add_shape(1, 0, 0, Inches(2.2), SH)
    band.fill.solid(); band.fill.fore_color.rgb = DARK; band.line.fill.background()
    add_text(s, Inches(0.2), Inches(0.6), Inches(2.0), Inches(0.6),
             f"STEP {num}", size=16, bold=True, color=YELLOW)
    add_text(s, Inches(0.2), Inches(1.3), Inches(2.0), Inches(4.0),
             "of 10", size=14, color=LIGHT)
    # Title
    add_text(s, Inches(2.6), Inches(0.5), Inches(10.3), Inches(0.9),
             title, size=32, bold=True, color=DARK)
    add_accent_bar(s, Inches(2.6), Inches(1.35), Inches(1.2), Inches(0.08), YELLOW)

    # What
    add_text(s, Inches(2.6), Inches(1.6), Inches(10.3), Inches(0.4),
             "WHAT", size=12, bold=True, color=YELLOW)
    add_text(s, Inches(2.6), Inches(2.0), Inches(10.3), Inches(1.6),
             what, size=15, color=DARK)
    # Why
    add_text(s, Inches(2.6), Inches(3.6), Inches(10.3), Inches(0.4),
             "WHY IT MATTERS", size=12, bold=True, color=YELLOW)
    add_text(s, Inches(2.6), Inches(4.0), Inches(10.3), Inches(1.4),
             why, size=15, color=DARK)
    # Done
    add_text(s, Inches(2.6), Inches(5.5), Inches(10.3), Inches(0.4),
             "DONE LOOKS LIKE", size=12, bold=True, color=YELLOW)
    add_text(s, Inches(2.6), Inches(5.9), Inches(10.3), Inches(1.4),
             done_text, size=15, color=DARK)


steps = [
    ("1", "Email delivery",
     "Connect SendGrid (or another email provider) so verification codes land in real inboxes, not just the server log. Replace the plain text email with a branded HTML template in Maverick colours. Send a one-time welcome email after a customer verifies.",
     "Right now the verification code only prints to the dev console — no real customer can sign up. Email is the front door to the whole product.",
     "A new customer receives a branded email within seconds, enters the code, and is logged in. A welcome email follows."),

    ("2", "Production database",
     "Apply the new schema (is_verified, email_verifications, unique index on transactions) to the production database before the first deploy. Decide what to do with users already in the dev database — typically mark them all as verified so nobody is locked out.",
     "Without the migration the live app will crash on sign-up. Without the backfill, existing testers get stuck at the verification screen.",
     "Production schema matches dev. Existing users can still log in. No 'column does not exist' errors in the logs."),

    ("3", "Account recovery",
     "Decide and build what happens when a customer loses access to their email or phone. Minimum: an admin can reset a user's email or phone from the admin dashboard, and the change is written to the audit log.",
     "Lost-access tickets are the #1 support cost for any loyalty programme. Having a clear path on day one prevents stuck customers and angry calls.",
     "Support can change a verified contact in under a minute. The change is logged and reversible."),

    ("4", "Legal & content pages",
     "Add Privacy Policy, Terms of Service, and a basic Help / FAQ page. Link them from the footer and from the 'by signing up you agree to…' line on the registration form. Get the wording reviewed by Maverick legal before publishing.",
     "POPIA (South Africa's privacy law) requires a published privacy policy before you collect personal data. Customers also need to know what they're agreeing to.",
     "Three pages live, linked from every page footer and the sign-up form. Legal has signed off on the wording."),

    ("5", "Branding polish",
     "Walk every page on mobile and desktop. Confirm logo, colours, fonts and tone match Maverick's brand kit. Check empty states look intentional, error messages are friendly, and loading states are visible.",
     "Customers judge trust in the first 5 seconds. A loyalty programme that looks like a developer demo loses sign-ups before the value is even explained.",
     "Every page passes a side-by-side visual review against Maverick's brand kit. No placeholder text remains."),

    ("6", "Real-device QA",
     "Test the full flow on a real Android phone, a real iPhone, and a desktop browser: sign up, receive the email, enter the code, upload a real MTN till slip or PDF invoice, see points, redeem a reward, log out, log back in, forgot password.",
     "Browser dev tools lie. Real cameras, real keyboards, and real network conditions catch the bugs that ship to thousands of customers.",
     "The end-to-end flow works on all three devices. Anything that broke has been fixed and re-tested."),

    ("7", "Admin & catalogue",
     "Create the first real super-admin account for Maverick and remove any test admins. Load the real rewards catalogue, real MTN packages, and real point-earning rules.",
     "Launching on placeholder rewards is worse than not launching — customers earn points they cannot spend on anything they want.",
     "Maverick's chosen admin can sign in. The rewards catalogue, packages, and earning rules match what marketing has approved."),

    ("8", "Deploy & domain",
     "Publish the app via Replit Deployments. If Maverick has a domain (e.g. loyalty.maverick.co.za), connect it and verify HTTPS. Set production secrets (DATABASE_URL, SESSION_SECRET, SendGrid key) in the deployment environment.",
     "Until it's deployed on a Maverick URL, it isn't a product — it's a prototype.",
     "The app is live on the chosen URL, HTTPS works, all secrets are set, and the production sign-up flow completes end-to-end."),

    ("9", "Monitoring & backups",
     "Confirm the deployment logs are accessible and you know how to read them. Set up a regular automated backup of the production Postgres database so you can recover from a bad deploy or data loss.",
     "Things will break in production — what matters is how quickly you see it and how cleanly you can roll back.",
     "Logs are reachable in two clicks. A nightly database backup exists and a test restore has been done."),

    ("10", "Go-live announcement",
     "Send the launch email to Maverick customers and update the Maverick website with a link to the loyalty portal. Coordinate with marketing on timing and messaging.",
     "A launch nobody knows about is the same as no launch.",
     "The announcement is sent. The Maverick website links to the loyalty portal. Sign-ups start arriving."),
]

for args in steps:
    step_slide(*args)

# Closing slide
s = prs.slides.add_slide(BLANK)
add_bg(s, DARK)
add_accent_bar(s, Inches(0.6), Inches(2.8), Inches(0.25), Inches(1.8))
add_text(s, Inches(1.0), Inches(2.7), Inches(11), Inches(1.0),
         "Ready when you are.", size=44, bold=True, color=WHITE)
add_text(s, Inches(1.0), Inches(3.7), Inches(11), Inches(0.7),
         "Pick a step and we'll start.", size=22, color=YELLOW)
add_text(s, Inches(1.0), Inches(4.6), Inches(11), Inches(0.5),
         "Recommended order: 1 → 2 → 7 → 4 → 8 → 3 → 5 → 6 → 9 → 10",
         size=14, color=LIGHT)

out = "Maverick_Loyalty_Launch_Checklist.pptx"
prs.save(out)
print("Saved:", out)
